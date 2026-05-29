import csv
import io
import os
import tempfile
from pathlib import Path

import mammoth
import openpyxl
import pandas as pd
from docx import Document


async def docx_to_text(content: bytes, filename: str) -> tuple[str, str]:
    document = Document(io.BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    stem = Path(filename).stem
    out_path = os.path.join(tempfile.mkdtemp(prefix="ff_out_"), f"{stem}.txt")
    Path(out_path).write_text(text, encoding="utf-8")
    return out_path, f"{stem}_converted.txt"


async def docx_to_html(content: bytes, filename: str) -> tuple[str, str]:
    result = mammoth.convert_to_html(io.BytesIO(content))
    stem = Path(filename).stem
    out_path = os.path.join(tempfile.mkdtemp(prefix="ff_out_"), f"{stem}.html")
    Path(out_path).write_text(result.value, encoding="utf-8")
    return out_path, f"{stem}_converted.html"


async def xlsx_to_csv(content: bytes, filename: str) -> tuple[str, str]:
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    stem = Path(filename).stem
    out_path = os.path.join(tempfile.mkdtemp(prefix="ff_out_"), f"{stem}.csv")
    with open(out_path, "w", encoding="utf-8", newline="") as csv_file:
        writer = csv.writer(csv_file)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(["" if cell is None else cell for cell in row])
    return out_path, f"{stem}_converted.csv"


async def csv_to_xlsx(content: bytes, filename: str) -> tuple[str, str]:
    dataframe = pd.read_csv(io.BytesIO(content))
    stem = Path(filename).stem
    out_path = os.path.join(tempfile.mkdtemp(prefix="ff_out_"), f"{stem}.xlsx")
    dataframe.to_excel(out_path, index=False)
    return out_path, f"{stem}_converted.xlsx"
